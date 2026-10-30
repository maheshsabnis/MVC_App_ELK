"""Create native, editable PowerPoint diagrams of the inspected solution."""

from pathlib import Path
from zipfile import ZipFile
import xml.etree.ElementTree as ET

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "elk-architecture-complete.pptx"
INK = "243746"
BLUE = "2764A5"
GREEN = "347B53"
ORANGE = "A25B12"
PURPLE = "7450A3"
GRAY = "66727C"
LIGHT = "F2F5F7"
prs = Presentation()
prs.slide_width = Inches(16)
prs.slide_height = Inches(9)
prs.core_properties.title = "MVC_App: actual telemetry and ELK architecture"
prs.core_properties.subject = "Editable diagrams based on four .NET 10 projects and observability configuration"
prs.core_properties.author = ""
prs.core_properties.keywords = "ASP.NET Core 10, Blazor, Razor Pages, OpenTelemetry, ELK, Azure Monitor"


def text(slide, x, y, w, h, content, size=16, color=INK, bold=False):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = Inches(0.03)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    for i, line in enumerate(content.split("\n")):
        paragraph = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        paragraph.text = line
        paragraph.font.name = "Aptos"
        paragraph.font.size = Pt(size)
        paragraph.font.bold = bold
        paragraph.font.color.rgb = RGBColor.from_string(color)
        paragraph.space_after = Pt(4)
    return shape


def slide(title, subtitle, notes):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    text(s, .4, .18, 15.2, .5, title, 27, bold=True)
    text(s, .4, .79, 15.2, .45, subtitle, 14, GRAY)
    text(s, .4, 8.67, 14.5, .2, "MVC_App solution | Current code/configuration, not a proposed deployment | Native editable shapes and connectors", 10, GRAY)
    text(s, 15.15, 8.64, .5, .25, str(len(prs.slides)), 11, GRAY)
    s.notes_slide.notes_text_frame.text = notes
    return s


def box(s, x, y, w, h, heading, body="", color=BLUE, size=16):
    shape = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(LIGHT)
    shape.line.color.rgb = RGBColor.from_string(color)
    shape.line.width = Pt(1.7)
    tf = shape.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Inches(.13)
    tf.margin_top = tf.margin_bottom = Inches(.08)
    for i, line in enumerate([heading] + body.split("\n") if body else [heading]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.name = "Aptos"
        p.font.size = Pt(size + 1 if i == 0 else size)
        p.font.bold = i == 0
        p.font.color.rgb = RGBColor.from_string(color if i == 0 else INK)
        p.space_after = Pt(4)
    return shape


def point(shape, side):
    x, y, w, h = shape.left, shape.top, shape.width, shape.height
    return {0: (x + w // 2, y), 1: (x, y + h // 2), 2: (x + w // 2, y + h), 3: (x + w, y + h // 2)}[side]


def arrow(s, a, b, start=3, end=1, color=GRAY, elbow=False, label=None, label_at=None):
    x1, y1 = point(a, start)
    x2, y2 = point(b, end)
    line = s.shapes.add_connector(MSO_CONNECTOR.ELBOW if elbow else MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    line.begin_connect(a, start)
    line.end_connect(b, end)
    line.line.color.rgb = RGBColor.from_string(color)
    line.line.width = Pt(1.8)
    head = OxmlElement("a:tailEnd")
    head.set("type", "triangle")
    head.set("w", "med")
    head.set("len", "med")
    line.line._get_or_add_ln().append(head)
    tree = s.shapes._spTree
    tree.remove(line._element)
    tree.insert(2, line._element)
    if label:
        assert label_at is not None
        text(s, *label_at, label, size=12, color=color)
    return line


def route(s, a, b, coordinates, start=2, end=2, color=GRAY):
    points = [point(a, start)] + [(Inches(x), Inches(y)) for x, y in coordinates] + [point(b, end)]
    for index, (first, last) in enumerate(zip(points, points[1:])):
        line = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, *first, *last)
        if index == 0:
            line.begin_connect(a, start)
        if index == len(points) - 2:
            line.end_connect(b, end)
            head = OxmlElement("a:tailEnd")
            head.set("type", "triangle")
            line.line._get_or_add_ln().append(head)
        line.line.color.rgb = RGBColor.from_string(color)
        line.line.width = Pt(1.8)
        tree = s.shapes._spTree
        tree.remove(line._element)
        tree.insert(2, line._element)


s = slide("1 | Whole solution: applications, databases and telemetry destinations",
          "Blue/gray: application traffic   •   Orange: OTLP/ELK telemetry   •   Purple: separate Azure Monitor export",
          "Sources: all four Program.cs files, both web TelemetryExtensions.cs files, PgSqlClient TelemetryBootstrap.cs, compose.yaml. "
          "MVC_App includes MVC, Razor Pages and JSON API endpoints. Blazor is interactive server, not WebAssembly. Both web apps use the Identity database but separate cookies. "
          "ApiClient has no OTel SDK. PgSqlClient exports to Azure Monitor, never to the local Collector. The ELK detail is expanded on slide 2. "
          "Browser users can also access MVC/Razor routes directly. Diagram connectors indicate dependency/request direction; return responses are omitted for clarity.")
browser = box(s, .45, 1.48, 2.15, 1.0, "Browser", "UI / forms / circuit", size=15)
blazor = box(s, 3.4, 1.48, 3.3, 1.0, "BlazorClient", "Interactive server + Identity", size=15)
mvc = box(s, 7.65, 1.48, 3.2, 1.0, "MVC_App", "MVC / Razor Pages / API", size=15)
company = box(s, 12.05, 1.35, 3.5, 1.15, "Company database", "SQL Server / configured SQL host", size=14)
identity = box(s, 12.05, 3.2, 3.5, 1.1, "MVC_AppIdentity", "Shared account database\nSeparate browser sessions", size=14)
api = box(s, .45, 3.2, 3.2, 1.1, "ApiClient console", "Scenario / load HTTP traffic\nNO OTel exporter", size=14)
collector = box(s, 7.65, 5.2, 3.2, 1.1, "OTel Collector", "Host loopback :4318\nLogs / traces / metrics", ORANGE, 14)
elk = box(s, 12.05, 5.2, 3.5, 1.1, "ELK stack", "Logstash • Elasticsearch\nKibana — detail on slide 2", ORANGE, 14)
pg = box(s, .45, 7.05, 3.1, .95, "PgSqlClient console", "Custom + Npgsql telemetry", PURPLE, 14)
pgdb = box(s, 5.05, 7.05, 3.1, .95, "PostgreSQL", "Independent Employee schema", GREEN, 14)
az = box(s, 10.1, 7.05, 5.45, .95, "Azure Monitor / Application Insights", "Separate destination — not connected to ELK", PURPLE, 14)
arrow(s, browser, blazor, label="HTTPS / SignalR", label_at=(2.65, 1.5, .7, .8))
arrow(s, blazor, mvc, label="API key +\ntrace context", label_at=(6.75, 1.4, .85, .85))
arrow(s, mvc, company, label="SQL", label_at=(11.1, 1.45, .7, .35))
arrow(s, mvc, identity, start=3, end=1, elbow=True)
route(s, blazor, identity, [(5.05, 2.9), (11.65, 2.9), (11.65, 3.75)], end=1)
text(s, 8.6, 2.65, 3.0, .28, "Identity SQL (both web apps)", 12)
route(s, api, mvc, [(4.05, 3.75), (4.05, 4.55), (7.2, 4.55), (7.2, 2.05)], start=3, end=1)
text(s, 4.2, 4.25, 2.7, .25, "HTTP workload → API", 12)
arrow(s, mvc, collector, start=2, end=0, color=ORANGE)
text(s, 9.35, 4.35, 1.45, .6, "OTLP\nHTTP/protobuf", 12, ORANGE)
route(s, blazor, collector, [(5.05, 5.75)], end=1, color=ORANGE)
text(s, 5.2, 5.43, 2.2, .3, "OTLP, all 3 signals", 12, ORANGE)
arrow(s, collector, elk, color=ORANGE, label="Split pipelines", label_at=(10.9, 5.3, 1.1, .65))
arrow(s, pg, pgdb, color=GREEN, label="Npgsql / TLS", label_at=(3.65, 7.08, 1.35, .35))
route(s, pg, az, [(2, 8.35), (12.82, 8.35)], color=PURPLE)
text(s, 4.6, 8.05, 5.5, .28, "Azure Monitor exporters: logs + traces + metrics", 12, PURPLE)
text(s, .45, 5.15, 3.5, 1.15, "Direct browser MVC/Razor access is also supported.\nApps and databases run outside this Compose file.", 14)

s = slide("2 | ELK deployment and signal routing",
          "compose.yaml + observability/otel-collector.yaml — only logs traverse Logstash; no Filebeat or APM Server",
          "Elasticsearch/Kibana/Logstash are pinned at 9.5.4; Collector Contrib is 0.161.0. Published ports are bound to 127.0.0.1. "
          "The Collector's receiver is HTTP only, not gRPC. All signal pipelines use memory_limiter then batch. "
          "Traces/metrics go directly to the elasticsearch exporter with require_data_stream=true; logs go to file/logs. "
          "The log file path is a host bind mount, not a Docker named log volume. logstash-data is a named volume for sincedb. "
          "No managed Elasticsearch named data volume is declared. Security is disabled for this local demonstration.")
apps = box(s, .45, 1.65, 3.0, 1.35, "Host web processes", "MVC_App\nMVC_App.BlazorClient\nOTLP HTTP/protobuf", size=14)
col = box(s, 4.5, 1.65, 4.25, 1.55, "Collector Contrib 0.161.0", "receiver: 0.0.0.0:4318\nmemory_limiter → batch\nContainer: 256 MB limit", ORANGE, 15)
es = box(s, 11.65, 1.65, 3.9, 1.55, "Elasticsearch 9.5.4", "OTel trace/metric data streams\nDaily mvcapp-logs-* indices\nHost 127.0.0.1:9200", BLUE, 14)
f = box(s, 4.5, 4.35, 4.25, 1.45, "Shared rotating JSONL", "Host: observability/otel-output\nContainer: /otel-output/logs.jsonl\n10 MB • 3 days • 3 backups", ORANGE, 14)
ls = box(s, 11.65, 4.35, 3.9, 1.45, "Logstash 9.5.4", "OTLP batch → individual events\n2 workers • batch size 8\n512 MB Java heap", ORANGE, 14)
kib = box(s, 11.65, 7.0, 3.9, .95, "Kibana 9.5.4", "Host 127.0.0.1:5601", BLUE, 15)
vol = box(s, 5.55, 7.0, 4.25, .95, "logstash-data named volume", "Persistent sincedb read offsets", GRAY, 14)
arrow(s, apps, col, label="Host :4318", label_at=(3.5, 2.05, .95, .65), color=ORANGE)
arrow(s, col, es, color=BLUE, label="traces + metrics\nElasticsearch exporter", label_at=(8.85, 1.7, 2.6, .9))
arrow(s, col, f, start=2, end=0, color=ORANGE, label="logs → file/logs", label_at=(6.8, 3.55, 2.3, .4))
arrow(s, f, ls, color=ORANGE, label="file input\njson codec", label_at=(9.0, 4.55, 2.3, .6))
arrow(s, ls, es, start=0, end=2, color=ORANGE, label="HTTP index writes", label_at=(13.85, 3.45, 1.65, .6))
route(s, es, kib, [(15.8, 2.42), (15.8, 7.48)], start=3, end=3, color=BLUE)
text(s, 13.05, 6.25, 2.45, .4, "Indexed data → searches / charts", 12, BLUE)
arrow(s, ls, vol, start=2, end=3, color=GRAY, elbow=True)
text(s, .45, 3.65, 3.2, 3.7, "Docker network: telemetry (bridge)\n\nInternal ES URL:\nhttp://elasticsearch:9200\n\nNo OTLP gRPC/4317 receiver.\nNo published Logstash input port.\nNo app or database services in Compose.", 15)
text(s, .45, 8.1, 14.9, .35, "Local-only: Elasticsearch auth/TLS disabled. Bounded queues + rotating files are not guaranteed delivery. Add explicit production security, storage and retention.", 12, GRAY)

s = slide("3 | Logstash transformation: OTLP envelopes → searchable records",
          "observability/logstash/pipeline/otel-logs.conf — preserving correlation without copying entire batches per log record",
          "One JSONL line is an ExportLogsServiceRequest envelope and can carry many records. file input separates lines; json decodes each document. "
          "Ruby walks resourceLogs, scopeLogs and logRecords and emits a fresh LogStash::Event per record via new_event_block. "
          "The AnyValue helper supports scalars, arrays and kvlists but not bytesValue. The timestamp conversion uses floating seconds. "
          "The wrapper is canceled after fan-out. Nested split filters would deep-copy large arrays and can exhaust memory. "
          "The http.route.matched rename avoids an Elasticsearch scalar/object conflict with the http.route string. "
          "No document_id is set, so replay can duplicate documents.")
a = box(s, .45, 1.7, 3.2, 1.55, "1. File input", "logs*.jsonl*\nstart_position: beginning\nPersisted sincedb offsets", ORANGE, 15)
b = box(s, 4.4, 1.7, 3.2, 1.55, "2. Decode envelope", "codec: json (not json_lines)\nresourceLogs[]\n→ scopeLogs[] → logRecords[]", ORANGE, 14)
c = box(s, 8.35, 1.7, 3.2, 1.55, "3. Ruby fan-out", "Decode AnyValue attributes\nFresh small event per record\nnew_event_block.call(out)", ORANGE, 14)
d = box(s, 12.3, 1.7, 3.2, 1.55, "4. Index events", "Cancel original envelope\nmvcapp-logs-YYYY.MM.dd\nhttp://elasticsearch:9200", BLUE, 14)
arrow(s, a, b, color=ORANGE)
arrow(s, b, c, color=ORANGE)
arrow(s, c, d, color=ORANGE)
fields = box(s, .45, 4.05, 7.1, 2.2, "Published log document fields", "@timestamp • message • severity_text / severity_number\ntrace_id • span_id • scope_name\nservice_name • deployment_environment\nresource_attributes • log_attributes", BLUE, 17)
normal = box(s, 8.1, 4.05, 7.4, 2.2, "Normalize legacy field spelling", "http.route.matched → http.route_matched\nhttp.route remains a scalar route-template string\nCurrent spelling wins if both are present\nPrevents object/scalar mapping rejection", GREEN, 17)
text(s, .55, 6.8, 7.0, 1.1, "Why not nested split filters?\nEach split can clone a large envelope. Fresh events keep the fan-out substantially smaller.", 17)
text(s, 8.2, 6.8, 7.1, 1.15, "Reliability boundary\nRemoving offsets can replay retained files. Rotation can delete unread logs. No deterministic document IDs are configured.", 17)

s = slide("4 | Request and error flow through ASP.NET Core 10",
          "HTTP trace propagation + structured logs + SQL spans; component events require additional Blazor observation points",
          "Sources: MVC_App Program.cs, middleware, DbCommandLoggingInterceptor, both web TelemetryExtensions, Blazor Program.cs and Telemetry folder. "
          "MVC request logging wraps the environment handler, application exception handler, status-code pages, HTTPS/static files, routing, auth and endpoints. "
          "API-key authentication is additive to cookie authentication. Models are validated explicitly for API problem consistency. "
          "Final completion logs and exception logs are distinct. Browser domain failures can redirect, giving a 302 final response alongside the original error event. "
          "Logs correlate with the ambient Activity; an interceptor callback is not guaranteed to execute inside the SQL span after provider completion. "
          "Blazor circuit events are not all normal HTTP requests, and /_blazor logging is skipped.")
ui = box(s, .45, 1.65, 2.65, 1.4, "Blazor component", "Typed API client\nApiLoggingHandler\nHttpClient span", BLUE, 14)
web = box(s, 3.75, 1.65, 3.05, 1.4, "MVC HTTP pipeline", "Request completion logger\nException / status handling\nRouting + API-key auth", BLUE, 14)
ctrl = box(s, 7.45, 1.65, 3.05, 1.4, "API controller", "Validate DTO / domain rules\nCompanyContext\nEF Core command execution", BLUE, 14)
sql = box(s, 11.15, 1.65, 4.4, 1.4, "SQL command + observation", "SqlClient span: db.statement\nInterceptor: operation / duration / outcome\nCompany and Identity contexts", BLUE, 14)
arrow(s, ui, web, label="HTTP +\ntraceparent", label_at=(3.12, 1.7, .6, .9))
arrow(s, web, ctrl)
arrow(s, ctrl, sql)
errors = box(s, .45, 4.0, 4.5, 2.1, "API error contract", "application/problem+json\nHTTP status + errorCategory / errorType\ntraceId for caller investigation\n400/404/409 → Warning; 5xx → Error", ORANGE, 16)
logs = box(s, 5.75, 4.0, 4.45, 2.1, "Common structured fields", "http.response.status_code\nerror.category / error.type\nevent.outcome • route / path\nTrace correlation is context-dependent", ORANGE, 16)
bound = box(s, 11.0, 4.0, 4.55, 2.1, "Blazor exception boundary", "LoggingErrorBoundaryLogger\nHandles escaping component failures\nApiErrorAlert shows handled API errors\nNo browser OTel SDK configured", ORANGE, 16)
text(s, .55, 6.75, 7.1, 1.35, "Count requests carefully\nFilter to request-completion scope_name. SQL, exception, outbound-call and framework events are not additional user requests.", 17)
text(s, 8.2, 6.75, 7.2, 1.35, "Keep classifications distinct\n503 has category service_unavailable. Use numeric status >= 500 for all 5xx, not only error.is_server_error.", 17)

s = slide("5 | Configuration, storage and operational boundaries",
          "The observability system needs its own monitoring, storage policy and access controls",
          "This deck documents existing settings, not recommended production defaults. JSON configuration is read explicitly for the web endpoint/name/enable switch. "
          "Standard SDK environment knobs such as OTEL_METRIC_EXPORT_INTERVAL are not proven to bind from appsettings JSON here; use real environment variables or explicit options. "
          "Collector send_batch_size is a send trigger, not a maximum; no send_batch_max_size is set. "
          "Log rotation is not Elasticsearch retention. Named sincedb storage is not indexed-data storage. "
          "No saved Kibana dashboards, Filebeat service, persistent Logstash queue or Collector storage extension is configured. "
          "Raw telemetry and credentials must not be copied to public articles. All diagram host/database labels are generic.")
box(s, .45, 1.6, 4.65, 2.45, "Application controls", "Telemetry:Enabled + OTLP endpoint gate\nDistinct service names for MVC and Blazor\nScopes carry the custom log fields\nProduction needs a supplied endpoint\nAppsettings / env / launch profile differ", BLUE, 16)
box(s, 5.65, 1.6, 4.65, 2.45, "Collector controls", "memory_limiter: 192 MiB, spike 48 MiB\nCheck: 1 second; batch timeout: 5 seconds\nsend_batch_size: 512 (not a hard maximum)\nElasticsearch queue: 256; max retries: 5\nNo disk-backed trace/metric queue", ORANGE, 16)
box(s, 10.85, 1.6, 4.65, 2.45, "Storage controls", "Logs: rotating host JSONL files\nOffsets: logstash-data named volume\nES: no managed named volume declared\nNo configured log-index lifecycle policy\nRotation != index retention", GREEN, 16)
box(s, .45, 4.65, 7.1, 2.7, "Current local-only assumptions", "Loopback ports: 4318 / 9200 / 5601\nElasticsearch auth and HTTP TLS disabled\nBlazor dev-only certificate-validation bypass\nSQL text, queries and user names may be logged\nFault-injection endpoints / load tests need isolation", ORANGE, 17)
box(s, 8.1, 4.65, 7.4, 2.7, "Before a production deployment", "Secure ingestion and dashboard access\nSecret rotation and external configuration\nExplicit durable storage, backups and retention\nSampling, cardinality limits and privacy review\nMonitor export failures, backlog, mappings and disk", BLUE, 17)
text(s, .55, 7.85, 14.9, .45, "Do not mistake these recommendations for already implemented features. This task documents the current solution; it does not deploy or reconfigure it.", 14, GRAY)

s = slide("6 | PostgreSQL client: a separate Azure Monitor telemetry branch",
          "MVC_App.PgSqlClient is part of the solution, but it does not export to the local ELK collector",
          "Sources: PgSqlClient Program.cs, ClientOptions.cs, EmployeeWorkload.cs, InteractiveSession.cs, EmployeeRepository.cs, AppDiagnostics.cs and TelemetryBootstrap.cs. "
          "The generic host registers a pooled NpgsqlDataSource with a shared logger factory and no parameter logging. "
          "Workload commands create business activities; repository wrappers create operation activities, record timing and errors; Npgsql adds command spans and metrics. "
          "Logs/traces/metrics use Azure Monitor exporters only when a connection string is configured; optional console exporters are independent. "
          "The Azure Monitor metric reader is minute-scale and the option for interval is for console export. Providers are force-flushed at shutdown. "
          "Gauges reflect the last table snapshot, not continuous polling. PostgreSQL SslMode.Require is encryption without VerifyFull validation.")
work = box(s, .45, 1.7, 3.4, 1.5, "Scenario / interactive session", "employee.scenario / employee.command.*\nIntentional failures are optional\nCancellation + final flush", PURPLE, 14)
repo = box(s, 4.7, 1.7, 4.0, 1.5, "EmployeeRepository", "Business operation activities\nCounters + duration histogram\nException / SQLSTATE recording", PURPLE, 15)
pgdb = box(s, 10.0, 1.7, 5.5, 1.5, "Npgsql → PostgreSQL", "Pooled connections; parameterized statements\nDriver spans and connection-pool metrics\nSeparate schema from MVC's SQL Server model", GREEN, 15)
arrow(s, work, repo, color=PURPLE)
arrow(s, repo, pgdb, color=GREEN, label="SQL / TLS", label_at=(8.8, 2.0, 1.1, .4))
diag = box(s, .45, 4.2, 4.4, 2.35, "AppDiagnostics + providers", "employee.operations / failures\nrows.affected / rows.read\noperation.duration (milliseconds)\nTable-row / active-row / salary gauges\nRuntime + process metrics", PURPLE, 15)
boot = box(s, 5.65, 4.2, 4.4, 2.35, "TelemetryBootstrap", "ILogger + scopes / parsed state\nActivitySource + Npgsql tracing\nMeter + Npgsql/runtime/process metrics\nConditional Azure Monitor exporters\nOptional console exporters", PURPLE, 15)
az = box(s, 11.1, 4.2, 4.4, 2.35, "Azure Monitor", "Application Insights destination\nLogs + traces + metrics\nConnection string supplied privately\nInvestigate using Azure/Kusto tooling\nNot mvcapp-logs-* in Kibana", PURPLE, 15)
arrow(s, diag, boot, color=PURPLE)
arrow(s, boot, az, color=PURPLE)
arrow(s, repo, boot, start=2, end=0, color=PURPLE)
text(s, .55, 7.1, 7.1, 1.05, "Short-lived process rule\nForceFlush on shutdown matters. A successful flush is not proof that every signal is already searchable in the backend.", 17)
text(s, 8.2, 7.1, 7.2, 1.05, "Two different KQLs\nAzure Kusto Query Language is not Kibana Query Language. Do not paste these systems' queries interchangeably.", 17)

s = slide("7 | From a user report to a useful Kibana investigation",
          "Log view: mvcapp-logs-* with @timestamp — trace/metric streams are separate and must be discovered",
          "Sources: Logstash output field assignments, request/error logging helpers and the API problem-details contract. "
          "Use service_name or resource_attributes.service.name for service identity. Use log_attributes prefixes for application scope fields. "
          "A trace_id from a returned problem document is a useful log query key when context was recorded. "
          "Do not claim all logs have a trace ID: startup/auth/component events may lack a current Activity. "
          "For request totals use the request logging middleware scope rather than all log documents. "
          "The smoke test checks host OTLP receipt only; it does not assert end-to-end Elasticsearch/Kibana ingestion. "
          "The PPTX contains native text, autoshapes and connectors. Edit nodes/connectors in PowerPoint; export a slide to PNG or SVG for Blogger.")
problem = box(s, .45, 1.7, 3.8, 1.4, "User / API error report", "Status + errorCategory\ntraceId (when available)\nTime and affected action", BLUE, 16)
search = box(s, 5.1, 1.7, 4.35, 1.4, "Kibana Discover", "Choose mvcapp-logs-*\nSet the time window\nSearch trace_id / service / category", BLUE, 16)
result = box(s, 10.3, 1.7, 5.2, 1.4, "Correlated evidence", "SQL event + request + error logs\nDistinguish exceptions from completions\nInspect traces/metrics in separate views", BLUE, 16)
arrow(s, problem, search, color=BLUE)
arrow(s, search, result, color=BLUE)
box(s, .45, 3.8, 7.1, 2.7, "Useful Kibana queries", 'service_name : "MVC_App.BlazorClient"\nlog_attributes.error.type : *\nlog_attributes.error.category : "conflict"\nlog_attributes.http.response.status_code >= 500\nlog_attributes.db.duration_ms > 100', BLUE, 16)
box(s, 8.1, 3.8, 7.4, 2.7, "Avoid misleading interpretations", "A bare error.type is the wrong field path\n409 is Warning, not necessarily Error\n503 != server_error category\n302 may follow a browser domain exception\nAll log documents != completed requests", ORANGE, 16)
text(s, .55, 7.05, 7.0, 1.1, "Validation ladder\nHost OTLP receipt → Collector export → file lag / Logstash → indexed documents → Kibana view/time range.", 17)
text(s, 8.2, 7.05, 7.2, 1.1, "Edit and publish\nAll boxes, labels and arrows are editable. Export a diagram image for Blogger; distribute the PPTX separately if desired.", 17)

s = slide("8 | Complete application and telemetry flow",
          "Application requests at the top • ELK signal routing in the middle • Independent PostgreSQL / Azure Monitor path below",
          "Editable diagram of the text architecture supplied for the Blogger article. Browser requests reach Blazor and MVC/Razor routes. "
          "Blazor calls the MVC API with HTTP trace context and accesses Identity SQL directly; MVC accesses Company SQL and Identity SQL. "
          "ApiClient generates HTTP workloads but has no console OTLP exporter. The middle band represents the combined telemetry emitted by the two web processes. "
          "OTLP HTTP/protobuf on port 4318 enters the Collector, where memory_limiter precedes batch. Traces and metrics go directly to Elasticsearch OTel data streams. "
          "Logs go through shared rotating JSONL and Logstash fan-out into mvcapp-logs-YYYY.MM.dd. Kibana queries Elasticsearch. "
          "PgSqlClient uses Npgsql for PostgreSQL and sends custom, Npgsql, runtime and process telemetry to Azure Monitor, not to the ELK collector. "
          "Sources: web Program.cs and TelemetryExtensions.cs, PgSqlClient TelemetryBootstrap.cs, compose.yaml, Collector YAML and Logstash pipeline. "
          "All nodes and arrows are native editable PowerPoint objects; export this slide as PNG for the slide-8 placeholder in article.html.")
browser = box(s, .45, 1.4, 1.9, .75, "Browser", size=15)
blazor = box(s, 3.8, 1.4, 2.8, .75, "Blazor server", size=15)
mvc = box(s, 8.0, 1.4, 2.6, .75, "MVC API", size=15)
company = box(s, 12.7, 1.4, 2.85, .75, "Company SQL", size=15)
identity = box(s, 12.7, 2.8, 2.85, .75, "Identity SQL", size=15)
api = box(s, .45, 2.85, 3.2, .8, "API console", "HTTP workload; no OTLP exporter", size=12)
arrow(s, browser, blazor, color=BLUE, label="UI requests", label_at=(2.45, 1.45, 1.3, .3))
arrow(s, blazor, mvc, color=BLUE, label="HTTP +\ntrace context", label_at=(6.75, 1.38, 1.15, .55))
arrow(s, mvc, company, color=BLUE, label="SQL", label_at=(11.35, 1.43, .7, .3))
route(s, browser, mvc, [(1.4, 2.5), (9.3, 2.5)], color=BLUE)
text(s, 1.65, 2.2, 3.2, .25, "Direct browser MVC / Razor routes", 11, BLUE)
route(s, blazor, identity, [(5.2, 3.175)], end=1, color=GRAY)
text(s, 9.6, 3.2, 2.9, .3, "Shared account database", 11, GRAY)
route(s, mvc, identity, [(11.8, 1.775), (11.8, 3.175)], start=3, end=1, color=GRAY)
route(s, api, mvc, [(7.35, 3.25), (7.35, 1.775)], start=3, end=1, color=BLUE)
text(s, 3.8, 3.45, 3.5, .28, "HTTP workload → MVC API", 11, BLUE)

signals = box(s, .45, 4.15, 3.2, 1.0, "MVC + Blazor OTel", "Logs • traces • metrics", ORANGE, 14)
collector = box(s, 4.7, 4.15, 3.35, 1.0, "OpenTelemetry Collector", "memory_limiter → batch", ORANGE, 14)
files = box(s, 4.7, 5.7, 3.35, .9, "Shared rotating JSONL", "Collector log-file output", ORANGE, 13)
logstash = box(s, 8.75, 5.7, 2.6, .9, "Logstash", "OTLP batch → events", ORANGE, 13)
es = box(s, 11.9, 3.95, 3.65, 2.35, "", color=BLUE)
es.fill.background()
text(s, 12.1, 4.02, 3.2, .3, "Elasticsearch", 17, BLUE, bold=True)
streams = box(s, 12.1, 4.5, 3.25, .6, "OTel data streams", size=14)
indices = box(s, 12.1, 5.45, 3.25, .6, "mvcapp-logs-YYYY.MM.dd", color=ORANGE, size=12)
kibana = box(s, 12.5, 6.6, 3.05, .55, "Kibana", size=14)
arrow(s, signals, collector, color=ORANGE, label="OTLP HTTP/\nprotobuf\n:4318", label_at=(3.7, 4.17, .95, .85))
arrow(s, collector, streams, color=BLUE, label="traces + metrics", label_at=(8.45, 4.35, 3.1, .3))
arrow(s, collector, files, start=2, end=0, color=ORANGE, label="logs", label_at=(6.55, 5.3, .8, .3))
arrow(s, files, logstash, color=ORANGE)
arrow(s, logstash, indices, color=ORANGE, elbow=True)
arrow(s, es, kibana, start=2, end=0, color=BLUE)
text(s, .55, 5.7, 3.5, .9, "Only logs pass through Logstash.\nTrace / metric export bypasses it.", 13, GRAY)

text(s, .45, 6.95, 11.4, .3, "Separate destination — not connected to the ELK collector", 15, PURPLE, bold=True)
pg = box(s, .45, 7.45, 3.2, .8, "PgSqlClient", "Independent console process", PURPLE, 13)
postgres = box(s, 5.25, 7.45, 3.2, .8, "PostgreSQL", "Database dependency", GREEN, 13)
azure = box(s, 11.9, 7.45, 3.65, .8, "Azure Monitor", "Logs • traces • metrics", PURPLE, 13)
arrow(s, pg, postgres, color=GREEN, label="Npgsql", label_at=(3.9, 7.48, 1.1, .3))
route(s, pg, azure, [(2.05, 8.48), (13.725, 8.48)], color=PURPLE)
text(s, 4.0, 8.17, 7.8, .28, "custom + Npgsql + runtime/process telemetry", 12, PURPLE)

prs.save(OUT)


def validate():
    loaded = Presentation(OUT)
    assert len(loaded.slides) == 8
    counts = {"slides": len(loaded.slides), "text_shapes": 0, "connectors": 0, "pictures": 0}
    for slide_number, item in enumerate(loaded.slides, 1):
        assert item.has_notes_slide, f"Missing source notes: slide {slide_number}"
        for shape in item.shapes:
            assert shape.left >= 0 and shape.top >= 0
            assert shape.left + shape.width <= loaded.slide_width + Inches(.02), f"Shape outside slide {slide_number}"
            assert shape.top + shape.height <= loaded.slide_height + Inches(.02), f"Shape outside slide {slide_number}"
            if shape.has_text_frame:
                counts["text_shapes"] += 1
            if shape.shape_type == MSO_SHAPE_TYPE.LINE:
                counts["connectors"] += 1
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                counts["pictures"] += 1
    assert counts["connectors"] > 30
    assert counts["pictures"] == 0, "Diagrams must not be flattened images"
    complete_flow = "\n".join(shape.text for shape in loaded.slides[7].shapes if shape.has_text_frame)
    for label in ("Browser", "Blazor server", "MVC API", "Company SQL", "Identity SQL", "API console",
                  "no OTLP exporter", "memory_limiter → batch", "Shared rotating JSONL", "Logstash",
                  "Elasticsearch", "OTel data streams", "mvcapp-logs-YYYY.MM.dd", "Kibana",
                  "PgSqlClient", "PostgreSQL", "Azure Monitor", "custom + Npgsql + runtime/process telemetry"):
        assert label in complete_flow, "Missing complete-flow component: " + label
    with ZipFile(OUT) as package:
        assert package.testzip() is None
        for name in package.namelist():
            if name.endswith((".xml", ".rels")):
                ET.fromstring(package.read(name))
    print("PASS editable PowerPoint:", counts)
    print(f"{OUT.name}: {OUT.stat().st_size:,} bytes")


if __name__ == "__main__":
    validate()
